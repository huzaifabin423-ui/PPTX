#!/usr/bin/env python3
"""
Bahria Town IT Security Analysis — PowerPoint generator
========================================================
Builds a clean, professional 16:9 presentation using python-pptx.

Student : Huzaifa Bin Mudassar
College : The Superior International College, Rahim Yar Khan
Company : Bahria Town Pvt. Ltd. — IT Department
Period  : April 2026 – June 2026

Run:
    pip install python-pptx
    python build_pptx.py
Output:
    Bahria_Town_Presentation.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --------------------------------------------------------------------------- #
#  Theme
# --------------------------------------------------------------------------- #
FONT = "Calibri"
PROJECT_NAME = "Bahria Town IT Security Analysis"

DARK_BG  = RGBColor(0x1A, 0x1F, 0x3A)   # dark blue background
CARD     = RGBColor(0x23, 0x2A, 0x52)   # card / panel
ROW1     = RGBColor(0x23, 0x2A, 0x52)   # table body row (odd)
ROW2     = RGBColor(0x2D, 0x36, 0x66)   # table body row (even)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
CYAN     = RGBColor(0x00, 0xD4, 0xFF)   # heading / accent
MUTED    = RGBColor(0x9A, 0xA3, 0xC7)   # footer text
DARKTEXT = RGBColor(0x0F, 0x14, 0x24)   # text on bright cells

RED    = RGBColor(0xFF, 0x5C, 0x5C)     # danger
ORANGE = RGBColor(0xFF, 0x9F, 0x40)     # high
YELLOW = RGBColor(0xFF, 0xD2, 0x4D)     # warning / medium
GREEN  = RGBColor(0x46, 0xD1, 0x7F)     # safe

SLIDE_W = 13.333
SLIDE_H = 7.5

# --------------------------------------------------------------------------- #
#  Presentation setup
# --------------------------------------------------------------------------- #
prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]


# --------------------------------------------------------------------------- #
#  Low-level helpers
# --------------------------------------------------------------------------- #
def new_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK_BG
    return slide


def no_line(shape):
    shape.line.fill.background()


def no_shadow(shape):
    shape.shadow.inherit = False


def rect(slide, left, top, width, height, color, line=None, line_w=1.0,
         shape_type=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape_type, Inches(left), Inches(top),
                                Inches(width), Inches(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        no_line(sp)
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    no_shadow(sp)
    return sp


def text_box(slide, left, top, width, height, text, *, size=18, color=WHITE,
             bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=None,
             font=FONT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top),
                                  Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    return tb


def run_in(p, text, *, size=18, color=WHITE, bold=False, font=FONT):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = font
    return r


# --------------------------------------------------------------------------- #
#  Slide chrome: header + footer
# --------------------------------------------------------------------------- #
def header(slide, title, number, subtitle=None):
    """Top accent bar, title, underline and footer."""
    rect(slide, 0, 0, SLIDE_W, 0.13, CYAN)                      # top edge bar
    text_box(slide, 0.6, 0.42, 12.1, 0.85, title,
             size=30, color=CYAN, bold=True)
    rect(slide, 0.62, 1.22, 3.4, 0.045, CYAN)                   # underline
    if subtitle:
        text_box(slide, 0.62, 1.30, 11.0, 0.4, subtitle,
                 size=14, color=MUTED, italic=True)
    footer(slide, number)


def footer(slide, number):
    rect(slide, 0.5, 6.95, 12.33, 0.014, MUTED)                 # thin divider
    text_box(slide, 0.5, 7.0, 8.0, 0.35, PROJECT_NAME,
             size=9, color=MUTED, align=PP_ALIGN.LEFT)
    text_box(slide, 10.83, 7.0, 2.0, 0.35, str(number),
             size=10, color=CYAN, bold=True, align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------- #
#  Bullet list
# --------------------------------------------------------------------------- #
def bullets(slide, items, *, left=0.75, top=1.6, width=11.9, height=5.1,
            size=19, gap=12, marker="▸"):
    """items: str  -> plain bullet
              (label, value) -> cyan bold label + white value"""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top),
                                  Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.12
        run_in(p, marker + "  ", size=size, color=CYAN, bold=True)
        if isinstance(item, tuple):
            label, value = item
            run_in(p, label + ":  ", size=size, color=CYAN, bold=True)
            run_in(p, value, size=size, color=WHITE)
        else:
            run_in(p, item, size=size, color=WHITE)
    return tb


# --------------------------------------------------------------------------- #
#  Card (rounded rectangle panel)
# --------------------------------------------------------------------------- #
def card(slide, left, top, w, h, title, body, *, title_size=18, body_size=15,
         title_color=CYAN, body_color=WHITE, anchor=MSO_ANCHOR.MIDDLE):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left),
                                Inches(top), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = CARD
    sp.line.color.rgb = CYAN
    sp.line.width = Pt(1)
    no_shadow(sp)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)
    p = tf.paragraphs[0]
    run_in(p, title, size=title_size, color=title_color, bold=True)
    body_items = body if isinstance(body, list) else [body]
    for j, line in enumerate(body_items):
        bp = tf.add_paragraph()
        bp.space_before = Pt(6)
        bp.line_spacing = 1.1
        if isinstance(body, list):
            run_in(bp, "•  ", size=body_size, color=CYAN, bold=True)
        run_in(bp, line, size=body_size, color=body_color)
    return sp


# --------------------------------------------------------------------------- #
#  Table
# --------------------------------------------------------------------------- #
def set_cell(cell, text, *, bold=False, color=WHITE, size=14,
             align=PP_ALIGN.LEFT, fill=None):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.14)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def build_table(slide, data, col_widths, left, top, *, header_size=15,
                body_size=14, row_height=0.55, color_rules=None):
    """color_rules: dict {col_index: func(text) -> (color, bold)}"""
    rows, cols = len(data), len(data[0])
    gtab = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                                  Inches(sum(col_widths)),
                                  Inches(row_height * rows))
    table = gtab.table
    table.first_row = False
    table.horz_banding = False
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for r in range(rows):
        table.rows[r].height = Inches(row_height)

    # header row
    for c in range(cols):
        set_cell(table.cell(0, c), data[0][c], bold=True, color=DARKTEXT,
                 size=header_size, fill=CYAN,
                 align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)

    # body rows
    for r in range(1, rows):
        fill = ROW1 if r % 2 == 1 else ROW2
        for c in range(cols):
            color, bold = WHITE, False
            if color_rules and c in color_rules:
                color, bold = color_rules[c](data[r][c])
            set_cell(table.cell(r, c), data[r][c], bold=bold, color=color,
                     size=body_size, fill=fill,
                     align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    return table


def risk_color(text):
    t = text.lower()
    if "very high" in t:
        return RED, True
    if "high" in t:
        return ORANGE, True
    if "medium" in t:
        return YELLOW, True
    return WHITE, False


def status_color(text):
    t = text.lower()
    if "not fixed" in t:
        return RED, True
    if "being fixed" in t or "in progress" in t:
        return YELLOW, True
    if "fixed" in t or "done" in t:
        return GREEN, True
    return WHITE, False


# =========================================================================== #
#  SLIDE 1 — Title
# =========================================================================== #
def slide_title():
    s = new_slide()
    rect(s, 0, 0, SLIDE_W, 0.18, CYAN)
    rect(s, 0, SLIDE_H - 0.18, SLIDE_W, 0.18, CYAN)
    text_box(s, 0.8, 0.95, 11.7, 0.5, "CYBER SECURITY PROJECT",
             size=18, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    text_box(s, 0.6, 2.05, 12.1, 1.9,
             "Bahria Town\nIT Security Analysis",
             size=48, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    rect(s, 5.17, 4.15, 3.0, 0.05, CYAN)                        # centred divider

    info = s.shapes.add_textbox(Inches(1.5), Inches(4.55),
                                Inches(10.33), Inches(2.2)).text_frame
    info.word_wrap = True
    lines = [
        ("Student:  ", "Huzaifa Bin Mudassar"),
        ("College:  ", "The Superior International College, Rahim Yar Khan"),
        ("Company:  ", "Bahria Town Pvt. Ltd. — IT Department"),
        ("Duration:  ", "April 2026  to  June 2026"),
    ]
    for i, (label, value) in enumerate(lines):
        p = info.paragraphs[0] if i == 0 else info.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(10)
        run_in(p, label, size=18, color=CYAN, bold=True)
        run_in(p, value, size=18, color=WHITE)
    footer(s, 1)


# =========================================================================== #
#  SLIDE 2 — Contents
# =========================================================================== #
def slide_contents():
    s = new_slide()
    header(s, "Contents", 2)
    toc = [
        (3, "About the Company"),
        (4, "IT Systems Used"),
        (5, "Network Setup"),
        (6, "Cyber Threats Found"),
        (7, "Risk Level Chart"),
        (8, "Security Problems Found"),
        (9, "Our Secure Website"),
        (10, "Security Features in Website"),
        (11, "Mobile and Cloud Risks"),
        (12, "What To Do When Hacked"),
        (13, "Data Protection"),
        (14, "Our Suggestions"),
        (15, "What I Learned"),
        (16, "Summary and Thank You"),
    ]

    def column(entries, left):
        tf = s.shapes.add_textbox(Inches(left), Inches(1.7),
                                  Inches(5.9), Inches(5.0)).text_frame
        tf.word_wrap = True
        for i, (num, title) in enumerate(entries):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(15)
            run_in(p, f"{num:02d}", size=18, color=CYAN, bold=True)
            run_in(p, "   " + title, size=18, color=WHITE)

    column(toc[:7], 1.1)
    column(toc[7:], 7.0)


# =========================================================================== #
#  SLIDE 3 — About the Company
# =========================================================================== #
def slide_about():
    s = new_slide()
    header(s, "About the Company", 3)
    bullets(s, [
        "Bahria Town is one of the largest real estate developers in Pakistan.",
        "Head offices are located in Rawalpindi and Islamabad.",
        "It runs a large, in-house IT department supporting all operations.",
        "The IT team manages websites, business software and networks.",
        "Thousands of staff and customers rely on these systems every day.",
    ], top=1.9, size=20, gap=16)


# =========================================================================== #
#  SLIDE 4 — IT Systems Used  (2x2 cards)
# =========================================================================== #
def slide_systems():
    s = new_slide()
    header(s, "IT Systems Used", 4)
    cards = [
        ("Computers", ["Dell servers", "Windows 10 / 11 workstations"]),
        ("Software",  ["ERP business system", "Microsoft 365 suite"]),
        ("Cloud",     ["Microsoft Azure", "Hosting & cloud storage"]),
        ("Security",  ["Kaspersky Antivirus", "CCTV surveillance"]),
    ]
    cw, ch, gx, gy = 5.75, 2.05, 0.45, 0.45
    x0, y0 = 0.69, 1.8
    for i, (title, body) in enumerate(cards):
        r, c = divmod(i, 2)
        card(s, x0 + c * (cw + gx), y0 + r * (ch + gy), cw, ch, title, body,
             title_size=20, body_size=15)


# =========================================================================== #
#  SLIDE 5 — Network Setup
# =========================================================================== #
def slide_network():
    s = new_slide()
    header(s, "Network Setup", 5)
    bullets(s, [
        "Built on Cisco routers and switches, secured by a Fortinet firewall.",
        "Uses a 3-layer design:  Core,  Distribution  and  Access.",
    ], top=1.6, size=18, gap=8, height=1.1)

    data = [
        ["Device", "Main Job in the Network"],
        ["Cisco Router", "Connects the office network to the internet"],
        ["Cisco Switch", "Links computers and devices inside the office"],
        ["Fortinet Firewall", "Blocks cyber attacks and filters traffic"],
        ["Core Layer", "Fast central backbone connecting everything"],
        ["Distribution Layer", "Controls routing and security between zones"],
        ["Access Layer", "Connects staff PCs, phones and printers"],
    ]
    build_table(s, data, [3.4, 7.6], left=1.17, top=2.75, row_height=0.52,
                header_size=15, body_size=14)


# =========================================================================== #
#  SLIDE 6 — Cyber Threats Found
# =========================================================================== #
def slide_threats():
    s = new_slide()
    header(s, "Cyber Threats Found", 6)
    data = [
        ["Threat", "Risk Level", "What It Means"],
        ["Phishing", "High Risk", "Fake emails trick staff into sharing passwords"],
        ["Ransomware", "Very High Risk", "Malware locks files and demands payment"],
        ["Insider Threats", "High Risk", "Staff misuse their access to company data"],
        ["Hacking Attempts", "Medium Risk", "Outsiders try to break into the systems"],
        ["Data Theft", "Very High Risk", "Customer or company data is stolen"],
    ]
    build_table(s, data, [2.7, 2.3, 6.5], left=0.92, top=1.95, row_height=0.74,
                header_size=15, body_size=14,
                color_rules={1: risk_color})


# =========================================================================== #
#  SLIDE 7 — Risk Level Chart  (5x5 matrix)
# =========================================================================== #
def slide_riskgrid():
    s = new_slide()
    header(s, "Risk Level Chart", 7,
           subtitle="Likelihood  ×  Impact  =  Risk Level")

    gl, gt, cw, ch = 2.55, 1.75, 1.16, 0.78
    threats = {(4, 5): "Ransomware", (3, 5): "Data Theft", (5, 3): "Phishing",
               (3, 4): "Insider", (2, 3): "Hacking"}

    for r in range(5):                       # top row = likelihood 5
        L = 5 - r
        for c in range(5):                   # left col = impact 1
            I = c + 1
            score = L * I
            color = GREEN if score <= 5 else (YELLOW if score <= 11 else RED)
            cell = rect(s, gl + c * cw, gt + r * ch, cw - 0.07, ch - 0.07,
                        color, line=DARK_BG, line_w=1.5)
            label = threats.get((L, I))
            if label:
                tf = cell.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                run_in(p, label, size=11, color=DARKTEXT, bold=True)

    # likelihood numbers (left)
    for r in range(5):
        text_box(s, gl - 0.45, gt + r * ch, 0.4, ch, str(5 - r),
                 size=13, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # impact numbers (bottom)
    for c in range(5):
        text_box(s, gl + c * cw, gt + 5 * ch + 0.02, cw - 0.07, 0.32,
                 str(c + 1), size=13, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER)

    # axis titles
    text_box(s, gl, gt + 5 * ch + 0.36, 5 * cw, 0.35,
             "Impact  →  (how bad the damage)",
             size=13, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    yaxis = text_box(s, 0.35, gt + 1.5, 2.4, 0.4,
                     "Likelihood  →  (how often)",
                     size=13, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    yaxis.rotation = 270

    # legend
    lx, ly = 9.35, 2.1
    text_box(s, lx, ly - 0.55, 3.4, 0.4, "Legend", size=16,
             color=CYAN, bold=True)
    legend = [(GREEN, "Safe  —  Low Risk"),
              (YELLOW, "Warning  —  Medium Risk"),
              (RED, "Danger  —  High Risk")]
    for i, (col, lab) in enumerate(legend):
        rect(s, lx, ly + i * 0.75, 0.42, 0.42, col)
        text_box(s, lx + 0.58, ly + i * 0.75 - 0.05, 3.15, 0.5, lab,
                 size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


# =========================================================================== #
#  SLIDE 8 — Security Problems Found
# =========================================================================== #
def slide_problems():
    s = new_slide()
    header(s, "Security Problems Found", 8)
    data = [
        ["Problem", "Severity Score", "Status"],
        ["Weak Passwords", "8 / 10", "Not Fixed"],
        ["Old Software", "7 / 10", "Being Fixed"],
        ["No Email Protection", "7 / 10", "Not Fixed"],
        ["No Staff Training", "9 / 10", "Not Fixed"],
    ]
    build_table(s, data, [5.0, 3.0, 3.0], left=1.17, top=2.0, row_height=0.78,
                header_size=15, body_size=15,
                color_rules={2: status_color})


# =========================================================================== #
#  SLIDE 9 — Our Secure Website
# =========================================================================== #
def slide_website():
    s = new_slide()
    header(s, "Our Secure Website", 9)
    bullets(s, [
        ("Project Name", "SecureShield"),
        "Built with Python Flask and a SQLite database.",
        "Secure login system with password protection.",
        "Tracks threats, incidents and risks in one place.",
        "Interactive dashboard with charts and reports.",
    ], left=0.75, top=1.9, width=7.2, size=18, gap=15)

    card(s, 8.35, 1.95, 4.3, 4.0, "Tech Stack",
         ["Python", "Flask (web framework)", "SQLite (database)",
          "HTML / CSS / Bootstrap", "Charts & Dashboard"],
         title_size=20, body_size=15, anchor=MSO_ANCHOR.TOP)


# =========================================================================== #
#  SLIDE 10 — Security Features in Website
# =========================================================================== #
def slide_features():
    s = new_slide()
    header(s, "Security Features in Our Website", 10)
    bullets(s, [
        "Passwords are encrypted (hashed) before they are stored.",
        "Login is protected — the account locks after 5 failed tries.",
        "Role-based access: users have different permission levels.",
        "Every action is recorded in audit logs for tracking.",
        "Protected against common attacks (SQL injection and XSS).",
    ], top=1.9, size=20, gap=16, marker="✔")


# =========================================================================== #
#  SLIDE 11 — Mobile and Cloud Risks
# =========================================================================== #
def slide_mobile_cloud():
    s = new_slide()
    header(s, "Mobile and Cloud Risks", 11)
    bullets(s, [
        "Staff use personal phones for work (BYOD risk).",
        "Some phones run old software that is not updated.",
        "People use public Wi-Fi, which is not safe for company data.",
        "Files are stored on the cloud without proper access controls.",
    ], top=1.95, size=20, gap=18)


# =========================================================================== #
#  SLIDE 12 — What To Do When Hacked  (5-step flow)
# =========================================================================== #
def slide_incident():
    s = new_slide()
    header(s, "What To Do When Hacked", 12,
           subtitle="Incident Response Plan — 5 Steps")
    steps = [
        ("1", "Identify", "Find the problem and detect the breach"),
        ("2", "Contain", "Stop it from spreading further"),
        ("3", "Eradicate", "Remove the threat completely"),
        ("4", "Recover", "Fix and restore the systems"),
        ("5", "Improve", "Learn lessons and improve defences"),
    ]
    cw, gap = 2.25, 0.18
    x0 = (SLIDE_W - (5 * cw + 4 * gap)) / 2
    ct, chh = 2.7, 2.5
    for i, (num, title, desc) in enumerate(steps):
        left = x0 + i * (cw + gap)
        box = slide_shapes_rounded(s, left, ct, cw, chh)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_top = Inches(0.55)
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run_in(p, title, size=16, color=CYAN, bold=True)
        dp = tf.add_paragraph()
        dp.alignment = PP_ALIGN.CENTER
        dp.space_before = Pt(8)
        dp.line_spacing = 1.1
        run_in(dp, desc, size=12.5, color=WHITE)

        # number badge straddling the top edge
        badge = rect(s, left + cw / 2 - 0.33, ct - 0.33, 0.66, 0.66, CYAN,
                     shape_type=MSO_SHAPE.OVAL)
        bf = badge.text_frame
        bf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = bf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        run_in(bp, num, size=22, color=DARKTEXT, bold=True)

        if i < 4:
            text_box(s, left + cw - 0.02, ct + chh / 2 - 0.35, gap + 0.1, 0.7,
                     "›", size=30, color=CYAN, bold=True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def slide_shapes_rounded(slide, left, top, w, h):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left),
                                Inches(top), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = CARD
    sp.line.color.rgb = CYAN
    sp.line.width = Pt(1)
    no_shadow(sp)
    return sp


# =========================================================================== #
#  SLIDE 13 — Data Protection
# =========================================================================== #
def slide_dataprotection():
    s = new_slide()
    header(s, "Data Protection", 13)
    bullets(s, [
        "Bahria Town stores customers' personal data (names, contacts, payments).",
        "This data must be kept safe, private and confidential.",
        "The company follows data protection rules and policies.",
        "Only authorised staff are allowed to access the data.",
    ], top=1.9, size=20, gap=16, height=3.2)

    callout = rect(s, 1.0, 5.55, 11.3, 0.85, CARD,
                   shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = callout.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run_in(p, "Core principle:  ", size=16, color=CYAN, bold=True)
    run_in(p, "Confidentiality  •  Integrity  •  Availability  "
              "(the CIA Triad)", size=16, color=WHITE)


# =========================================================================== #
#  SLIDE 14 — Our Suggestions  (roadmap bands)
# =========================================================================== #
def slide_suggestions():
    s = new_slide()
    header(s, "Our Suggestions", 14)
    phases = [
        ("RIGHT NOW", "Turn on 2-step login (2FA) for all staff accounts"),
        ("THIS MONTH", "Train all staff about cyber safety and awareness"),
        ("THIS YEAR", "Set up 24/7 security monitoring (SOC)"),
        ("FUTURE", "Follow international security standards (ISO 27001)"),
    ]
    bl, bw, bh, gap, y0 = 0.9, 11.5, 0.95, 0.28, 1.8
    for i, (phase, action) in enumerate(phases):
        y = y0 + i * (bh + gap)
        rect(s, bl, y, bw, bh, CARD, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        badge = rect(s, bl + 0.18, y + 0.16, 2.5, bh - 0.32, CYAN,
                     shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        bf = badge.text_frame
        bf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = bf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        run_in(bp, phase, size=15, color=DARKTEXT, bold=True)
        text_box(s, bl + 2.95, y, bw - 3.1, bh, action,
                 size=17, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


# =========================================================================== #
#  SLIDE 15 — What I Learned
# =========================================================================== #
def slide_learned():
    s = new_slide()
    header(s, "What I Learned", 15)
    bullets(s, [
        "Completed an 8-week internship at the Bahria Town IT Department.",
        "Learned how real companies protect their data and systems.",
        "Built a working security web application (SecureShield).",
        "Wrote a full security analysis report.",
        "Gained skills in network setup and threat analysis.",
    ], top=1.9, size=20, gap=16)


# =========================================================================== #
#  SLIDE 16 — Summary and Thank You
# =========================================================================== #
def slide_summary():
    s = new_slide()
    header(s, "Summary and Thank You", 16)
    bullets(s, [
        "Found many real security problems across the IT systems.",
        "Built SecureShield to help track and fix those problems.",
        "Delivered a full report with clear, practical recommendations.",
    ], top=1.75, size=19, gap=14, height=2.6)

    rect(s, 5.17, 4.45, 3.0, 0.05, CYAN)
    text_box(s, 1.0, 4.65, 11.33, 0.9, "Thank You",
             size=40, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    text_box(s, 1.3, 5.75, 10.73, 0.9,
             "Special thanks to Bahria Town Pvt. Ltd. (IT Department) "
             "and The Superior International College, Rahim Yar Khan.",
             size=15, color=WHITE, align=PP_ALIGN.CENTER)


# --------------------------------------------------------------------------- #
#  Build
# --------------------------------------------------------------------------- #
def main():
    slide_title()
    slide_contents()
    slide_about()
    slide_systems()
    slide_network()
    slide_threats()
    slide_riskgrid()
    slide_problems()
    slide_website()
    slide_features()
    slide_mobile_cloud()
    slide_incident()
    slide_dataprotection()
    slide_suggestions()
    slide_learned()
    slide_summary()

    out = "Bahria_Town_Presentation.pptx"
    prs.save(out)
    print(f"Saved {out} with {len(prs.slides._sldIdLst)} slides.")


if __name__ == "__main__":
    main()
